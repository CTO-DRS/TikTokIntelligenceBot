"""
Excel report (openpyxl) — Overview / Videos / Analytics / Growth sheets with
native Excel charts (LineChart for followers, BarChart for top videos,
PieChart for the content mix).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from bot.utils.i18n import t


class ExcelReport:
    def __init__(self, lang: str = "ar"):
        self.lang = lang if lang in {"ar", "en"} else "ar"

    # helpers -----------------------------------------------------
    def _hdr_style(self, ws, row: int, cols: int) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C2185B")
            cell.alignment = Alignment(horizontal="center")

    def build(self, profile: dict, analytics: dict, snapshots: list, out_path: Path) -> Path:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference

        wb = Workbook()

        # ── Sheet 1: overview ────────────────────────────────
        ws = wb.active
        ws.title = "Overview"
        ws.append(["TikTok Intelligence Report", f"@{profile.get('username')}"])
        ws.append(["generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
        ws.append([])
        headers = [t(self.lang, k) for k in
                   ("f_followers", "f_following", "f_likes", "f_videos",
                    "f_avg_views", "f_engagement", "security_score")]
        values = [
            profile.get("followers") or 0,
            profile.get("following") or 0,
            profile.get("likes") or 0,
            profile.get("videos_count") or 0,
            round(analytics.get("avg_views") or 0, 1),
            round(analytics.get("engagement_rate") or 0, 3),
            (profile.get("security") or {}).get("score"),
        ]
        ws.append(headers)
        ws.append(values)
        self._hdr_style(ws, 4, len(headers))
        for col, w in zip("ABCDEFG", (18, 16, 14, 12, 14, 14, 14)):
            ws.column_dimensions[col].width = w

        # ── Sheet 2: videos ──────────────────────────────────
        wsv = wb.create_sheet("Videos")
        cols = ["#", "id", "datetime", "duration_s", t(self.lang, "view_video"),
                "likes", "comments", "shares", "bookmarks", "hashtags", "music"]
        wsv.append(cols)
        self._hdr_style(wsv, 1, len(cols))
        for i, v in enumerate((profile.get("videos") or [])[:30], 1):
            ts = v.get("create_time")
            dt = datetime.utcfromtimestamp(int(ts)).strftime("%Y-%m-%d %H:%M") if ts else ""
            wsv.append([i, str(v.get("id")), dt, v.get("duration"),
                        v.get("views"), v.get("likes"), v.get("comments"),
                        v.get("shares"), v.get("bookmarks"),
                        " ".join("#" + h for h in (v.get("hashtags") or [])),
                        f"{v.get('music_title') or ''} — {v.get('music_author') or ''}"])
        for col, w in zip("ABCDEFGHIJK", (4, 20, 17, 10, 12, 10, 10, 9, 10, 26, 30)):
            wsv.column_dimensions[col].width = w

        n_videos = min(10, max(1, wsv.max_row - 1))
        bar = BarChart()
        bar.title = t(self.lang, "view_video")
        bar.height, bar.width = 7, 16
        data_ref = Reference(wsv, min_col=5, min_row=1, max_row=n_videos + 1)
        cats_ref = Reference(wsv, min_col=1, min_row=2, max_row=n_videos + 1)
        bar.add_data(data_ref, titles_from_data=True)
        bar.set_categories(cats_ref)
        wsv.add_chart(bar, "M2")

        # ── Sheet 3: analytics ───────────────────────────────
        wsa = wb.create_sheet("Analytics")
        wsa.append(["metric", "value"])
        self._hdr_style(wsa, 1, 2)
        for key, val in {
            "videos_analyzed": analytics.get("videos_analyzed"),
            "avg_views": analytics.get("avg_views"),
            "avg_likes": analytics.get("avg_likes"),
            "avg_comments": analytics.get("avg_comments"),
            "avg_shares": analytics.get("avg_shares"),
            "engagement_rate_%": analytics.get("engagement_rate"),
            "er_basis": analytics.get("er_basis"),
        }.items():
            wsa.append([key, val])
        row = wsa.max_row + 2
        wsa.append([t(self.lang, "f_categories"), "count"])
        self._hdr_style(wsa, row, 2)
        for cat, count in (analytics.get("categories") or {}).items():
            wsa.append([cat, count])

        pie = PieChart()
        pie.title = t(self.lang, "f_categories")
        pie.height, pie.width = 7, 10
        if (analytics.get("categories") or {}).items():
            cat_count = len(list(analytics["categories"].items()))
            pie.add_data(Reference(wsa, min_col=2, min_row=row + 1, max_row=row + cat_count),
                         titles_from_data=False)
            pie.set_categories(Reference(wsa, min_col=1, min_row=row + 1, max_row=row + cat_count))
        wsa.add_chart(pie, "D2")

        # ── Sheet 4: growth (snapshots) ──────────────────────
        wsg = wb.create_sheet("Growth")
        wsg.append(["date", t(self.lang, "f_followers"), "following", "likes", "videos"])
        self._hdr_style(wsg, 1, 5)
        for s in snapshots:
            wsg.append([s.captured_at.strftime("%Y-%m-%d %H:%M") if s.captured_at else "",
                        s.followers, s.following, s.likes, s.video_count])
        wsg.column_dimensions["A"].width = 18
        if len(snapshots) >= 2:
            line = LineChart()
            line.title = t(self.lang, "f_growth")
            line.height, line.width = 8, 16
            line.add_data(Reference(wsg, min_col=2, min_row=1, max_row=len(snapshots) + 1),
                          titles_from_data=True)
            line.set_categories(Reference(wsg, min_col=1, min_row=2, max_row=len(snapshots) + 1))
            wsg.add_chart(line, "F2")

        wb.save(out_path)
        return out_path
